"""Gerçek duman testi: tiny modeliyle CPU üzerinde tüm hat (ses çözme, yazıya dökme, hizalama, dışa aktarma).

İlk çalıştırmada tiny modeli ve hizalama modelini internetten indirir. Atlamak için: pytest -m "not smoke"
"""
import io
import json
import re
import time

import pytest
from docx import Document

from app.util import tr_lower
from tests.conftest import ROOT_DIR

SAMPLE = ROOT_DIR / "tests" / "data" / "ornek_tr.wav"
TRUTH = (ROOT_DIR / "tests" / "data" / "ornek_tr.txt").read_text(encoding="utf-8")

pytestmark = pytest.mark.smoke


def norm_words(text: str) -> list[str]:
    return re.sub(r"[^\w\s]", " ", tr_lower(text)).split()


def wer(ref: list[str], hyp: list[str]) -> float:
    d = list(range(len(hyp) + 1))
    for i, r in enumerate(ref, 1):
        prev, d[0] = d[0], i
        for j, h in enumerate(hyp, 1):
            prev, d[j] = d[j], min(d[j] + 1, d[j - 1] + 1, prev + (r != h))
    return d[len(hyp)] / max(1, len(ref))


@pytest.fixture()
def real_client(tmp_path, monkeypatch):
    monkeypatch.setenv("DESIFRE_DATA", str(tmp_path / "data"))
    monkeypatch.delenv("DESIFRE_ENGINE", raising=False)
    from fastapi.testclient import TestClient
    from app.main import create_app
    with TestClient(create_app()) as c:
        c.put("/api/settings", json={"device": "cpu", "model": "tiny", "align": True})
        yield c


def test_full_pipeline_with_tiny_model_on_cpu(real_client):
    c = real_client
    r = c.post("/api/jobs", params={"filename": "örnek kayıt.wav", "options": json.dumps({"model": "tiny", "language": "tr"})},
               content=SAMPLE.read_bytes())
    jid = r.json()["id"]
    end = time.time() + 900
    job = {}
    while time.time() < end:
        job = c.get(f"/api/jobs/{jid}").json()
        if job["status"] in ("tamamlandı", "hata", "iptal"):
            break
        time.sleep(1)
    assert job["status"] == "tamamlandı", job.get("error")
    assert not job["notice"], job["notice"]

    tr = c.get(f"/api/transcripts/{jid}").json()
    assert tr["language"] == "tr" and tr["model"] == "tiny"
    assert 20 < tr["duration"] < 30
    segs = tr["segments"]
    assert len(segs) >= 2, "Cümle bazlı bölümlere ayrılmalı"
    hyp = norm_words(" ".join(s["text"] for s in segs))
    error = wer(norm_words(TRUTH), hyp)
    print(f"\nWER (tiny, CPU): {error:.2f}\n" + "\n".join(f"  [{s['start']:.1f}-{s['end']:.1f}] {s['text']}" for s in segs))
    assert error < 0.6

    # zamanlar mantıklı ve sıralı
    prev_end = 0.0
    for s in segs:
        assert s["start"] >= prev_end - 0.5 and s["end"] > s["start"]
        assert [w["word"] for w in s["words"]] == s["text"].split()
        assert all(w["start"] <= w["end"] for w in s["words"])
        prev_end = s["end"]
    assert segs[-1]["end"] <= tr["duration"] + 0.5

    exp = lambda fmt, **p: c.get(f"/api/transcripts/{jid}/export/{fmt}", params=p)
    srt = exp("srt").content.decode("utf-8-sig")
    times = re.findall(r"(\d\d):(\d\d):(\d\d),(\d\d\d) --> (\d\d):(\d\d):(\d\d),(\d\d\d)", srt)
    assert len(times) >= 3
    secs = [(int(a) * 3600 + int(b) * 60 + int(c_) + int(d) / 1000, int(e) * 3600 + int(f) * 60 + int(g) + int(h) / 1000) for a, b, c_, d, e, f, g, h in times]
    assert all(s < e for s, e in secs) and all(secs[i][1] <= secs[i + 1][0] + 1e-6 for i in range(len(secs) - 1))
    assert all(len(line) <= 42 for line in srt.splitlines() if "-->" not in line and not line.strip().isdigit())
    assert exp("vtt").text.startswith("WEBVTT")
    assert "başlangıç;bitiş;konuşmacı;metin" in exp("csv").content.decode("utf-8-sig")
    assert json.loads(exp("json").content)["segments"][0]["words"]
    assert exp("txt", timestamps="0", speakers="0").content.decode("utf-8-sig").strip()
    doc = Document(io.BytesIO(exp("docx").content))
    assert doc.paragraphs[0].text == "örnek kayıt" and any(p.text.strip() for p in doc.paragraphs[2:])
    for fmt in ("txt", "srt", "csv"):
        assert exp(fmt).content.startswith(b"\xef\xbb\xbf")


def test_diarize_without_token_is_skipped_with_notice(real_client):
    c = real_client
    r = c.post("/api/jobs", params={"filename": "a.wav", "options": json.dumps({"model": "tiny", "diarize": True})},
               content=SAMPLE.read_bytes())
    jid = r.json()["id"]
    end = time.time() + 600
    while time.time() < end:
        job = c.get(f"/api/jobs/{jid}").json()
        if job["status"] in ("tamamlandı", "hata"):
            break
        time.sleep(1)
    assert job["status"] == "tamamlandı", job.get("error")
    assert "belirteci" in job["notice"]
    tr = c.get(f"/api/transcripts/{jid}").json()
    assert tr["segments"] and tr["speakers"] == {} and all(s["speaker"] is None for s in tr["segments"])
    assert "belirteci" in tr["notice"]

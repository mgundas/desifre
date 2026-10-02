"""DOCX dışa aktarma."""
from __future__ import annotations

import io
import time

from docx import Document
from docx.shared import Pt, RGBColor

from ..languages import language_name
from ..util import fmt_clock, speaker_label
from .txt import group_paragraphs

MONTHS = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def tr_date(ts: float) -> str:
    t = time.localtime(ts)
    return f"{t.tm_mday} {MONTHS[t.tm_mon - 1]} {t.tm_year} {t.tm_hour:02d}:{t.tm_min:02d}"


def _rgb(hexstr: str | None) -> RGBColor | None:
    try:
        h = (hexstr or "").lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except (ValueError, IndexError):
        return None


def to_docx(tr: dict, timestamps: bool = True, show_speakers: bool = True) -> bytes:
    doc = Document()
    doc.add_heading(tr["title"], level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Tarih: {tr_date(tr['created_at'])}\n")
    meta.add_run(f"Süre: {fmt_clock(tr['duration'])}\n")
    meta.add_run(f"Dil: {language_name(tr['language'])}")
    for r in meta.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    speakers = tr.get("speakers") or {}
    for p in group_paragraphs(tr["segments"]):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(8)
        if timestamps:
            r = para.add_run(f"[{fmt_clock(p['start'])}] ")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        label = speaker_label(p["speaker"], speakers) if show_speakers else ""
        if label:
            r = para.add_run(f"{label}: ")
            r.bold = True
            color = _rgb((speakers.get(p["speaker"]) or {}).get("color"))
            if color is not None:
                r.font.color.rgb = color
        para.add_run(p["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
